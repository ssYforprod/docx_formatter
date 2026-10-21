from __future__ import annotations

from pathlib import Path

from docx import Document

from docx_formatter.core.pipeline import Formatter
from docx_formatter.core.profile import Profile


def _make_table_docx(tmp_path: Path) -> Path:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).paragraphs[0].text = "Параметр"
    table.cell(0, 1).paragraphs[0].text = "Значение"
    table.cell(1, 0).paragraphs[0].text = "Температура"
    table.cell(1, 1).paragraphs[0].text = "20"
    path = tmp_path / "input.docx"
    document.save(path)
    return path


def test_table_cells_get_profile_font_and_header_bold(tmp_path: Path, gost_profile: Profile) -> None:
    input_path = _make_table_docx(tmp_path)

    out_path, report = Formatter(gost_profile).run(input_path)

    document = Document(out_path)
    table = document.tables[0]
    tables_cfg = gost_profile.tables

    header_run = table.cell(0, 0).paragraphs[0].runs[0]
    assert header_run.bold is tables_cfg.header_bold
    assert header_run.font.size.pt == tables_cfg.size_pt

    body_run = table.cell(1, 0).paragraphs[0].runs[0]
    assert body_run.font.size.pt == tables_cfg.size_pt

    assert any(c.rule_id == "tables" for c in report.changes)


def test_idempotent_with_table(tmp_path: Path, gost_profile: Profile) -> None:
    input_path = _make_table_docx(tmp_path)

    first_out, first_report = Formatter(gost_profile).run(input_path)
    assert any(c.rule_id == "tables" for c in first_report.changes)

    _, second_report = Formatter(gost_profile).run(first_out)

    assert second_report.changes == []
