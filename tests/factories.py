from __future__ import annotations

from pathlib import Path

from docx import Document


def make_docx(tmp_path: Path, blocks: list[tuple]) -> Path:
    """Build a minimal .docx from blocks, for use as test input.

    Supported blocks: ("p", text), ("table", rows, cols).
    """
    document = Document()
    for block in blocks:
        kind = block[0]
        if kind == "p":
            _, text = block
            document.add_paragraph(text)
        elif kind == "table":
            _, rows, cols = block
            document.add_table(rows=rows, cols=cols)
        else:
            raise ValueError(f"unknown block kind: {kind!r}")

    path = tmp_path / "input.docx"
    document.save(path)
    return path
