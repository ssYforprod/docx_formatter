from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement

from docx_formatter.core.captions import ParsedCaption, paragraph_has_drawing, parse_caption


def test_parses_figure_caption_with_em_dash() -> None:
    assert parse_caption("Рисунок 1 — Схема установки") == ParsedCaption(
        number="1", description="Схема установки"
    )


def test_parses_table_caption_with_hyphen() -> None:
    assert parse_caption("Таблица 2.1 - Результаты измерений") == ParsedCaption(
        number="2.1", description="Результаты измерений"
    )


def test_parses_abbreviated_prefix() -> None:
    assert parse_caption("Рис. 3 Внешний вид") == ParsedCaption(number="3", description="Внешний вид")


def test_rejects_non_caption_text() -> None:
    assert parse_caption("Это обычный абзац текста.") is None


def test_paragraph_with_inline_drawing_is_detected() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("рисунок здесь")
    run._element.append(OxmlElement("w:drawing"))

    assert paragraph_has_drawing(paragraph) is True


def test_paragraph_without_drawing_is_not_detected() -> None:
    document = Document()
    paragraph = document.add_paragraph("обычный текст")

    assert paragraph_has_drawing(paragraph) is False
