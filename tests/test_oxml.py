from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from factories import make_docx

from docx_formatter.core.oxml import iter_block_items, set_run_font


def _rfonts_element(run):
    rPr = run._element.find(qn("w:rPr"))
    assert rPr is not None, "ожидался w:rPr после set_run_font"
    rFonts = rPr.find(qn("w:rFonts"))
    assert rFonts is not None, "ожидался w:rFonts после set_run_font"
    return rFonts


def test_set_run_font_sets_all_four_script_attributes() -> None:
    document = Document()
    run = document.add_paragraph("Привет").runs[0]

    set_run_font(run, "Times New Roman")

    rFonts = _rfonts_element(run)
    # run.font.name пишет только ascii/hAnsi (разведка) — без cs/eastAsia
    # кириллица и complex-script текст не подхватывают новый шрифт.
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        assert rFonts.get(qn(attr)) == "Times New Roman"


def test_set_run_font_removes_theme_attributes() -> None:
    document = Document()
    run = document.add_paragraph("Привет").runs[0]
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:asciiTheme"), "minorHAnsi")
    rFonts.set(qn("w:cstheme"), "minorBidi")

    set_run_font(run, "Times New Roman")

    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        assert rFonts.get(qn(attr)) is None


def test_iter_block_items_preserves_document_order(tmp_path: Path) -> None:
    path = make_docx(
        tmp_path,
        blocks=[("p", "до таблицы"), ("table", 2, 3), ("p", "после таблицы")],
    )
    document = Document(path)

    items = list(iter_block_items(document))

    assert [type(item) for item in items] == [Paragraph, Table, Paragraph]
    assert items[0].text == "до таблицы"
    assert len(items[1].rows) == 2
    assert len(items[1].columns) == 3
    assert items[2].text == "после таблицы"


def test_iter_block_items_skips_section_properties(tmp_path: Path) -> None:
    # document.element.body всегда заканчивается w:sectPr — разведка
    # подтвердила, что без фильтрации по тегу он попадёт в обход как
    # «третий блок».
    path = make_docx(tmp_path, blocks=[("p", "единственный абзац")])
    document = Document(path)

    items = list(iter_block_items(document))

    assert len(items) == 1
