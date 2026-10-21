from __future__ import annotations

from dataclasses import replace

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from docx_formatter.core.detect import detect_heading
from docx_formatter.core.profile import HeadingsProfile

_PROFILE = HeadingsProfile(
    trust_styles_only=False,
    detect_threshold=0.75,
    levels=(),
    keywords_level1=("ВВЕДЕНИЕ", "ЗАКЛЮЧЕНИЕ"),
)
_BODY_SIZE_PT = 14.0


def test_english_heading_style_is_full_confidence() -> None:
    document = Document()
    paragraph = document.add_paragraph("Обзор", style="Heading 1")

    match = detect_heading(paragraph, _PROFILE, _BODY_SIZE_PT)

    assert match is not None
    assert match.by_style is True
    assert match.confidence == 1.0
    assert match.level == 1


def test_russian_heading_style_name_is_recognised() -> None:
    document = Document()
    document.styles.add_style("Заголовок 2", WD_STYLE_TYPE.PARAGRAPH)
    paragraph = document.add_paragraph("Методы исследования", style="Заголовок 2")

    match = detect_heading(paragraph, _PROFILE, _BODY_SIZE_PT)

    assert match is not None
    assert match.by_style is True
    assert match.level == 2


def test_manual_numbered_heading_scores_above_format_threshold() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("1 ВВЕДЕНИЕ")
    run.bold = True
    run.font.size = Pt(16)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    match = detect_heading(paragraph, _PROFILE, _BODY_SIZE_PT)

    assert match is not None
    assert match.by_style is False
    assert match.level == 1
    assert match.confidence >= 0.75


def test_ambiguous_bold_sentence_scores_in_warning_band() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Общие положения")
    run.bold = True

    match = detect_heading(paragraph, _PROFILE, _BODY_SIZE_PT)

    assert match is not None
    assert 0.4 <= match.confidence < 0.75


def test_ordinary_sentence_scores_below_warning_floor() -> None:
    document = Document()
    paragraph = document.add_paragraph(
        "Это обычное предложение из отчёта, которое заканчивается точкой."
    )

    match = detect_heading(paragraph, _PROFILE, _BODY_SIZE_PT)

    assert match is not None
    assert match.confidence < 0.4


def test_trust_styles_only_ignores_manual_formatting() -> None:
    profile = replace(_PROFILE, trust_styles_only=True)
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("1 ВВЕДЕНИЕ")
    run.bold = True
    run.font.size = Pt(16)

    match = detect_heading(paragraph, profile, _BODY_SIZE_PT)

    assert match is not None
    assert match.confidence == 0.0


def test_empty_paragraph_is_not_a_heading() -> None:
    document = Document()
    paragraph = document.add_paragraph("   ")

    assert detect_heading(paragraph, _PROFILE, _BODY_SIZE_PT) is None
