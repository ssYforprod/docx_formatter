from __future__ import annotations

from dataclasses import replace
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

from docx_formatter.core.pipeline import Formatter
from docx_formatter.core.profile import Profile


def _add_figure(document: Document) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    run._element.append(OxmlElement("w:drawing"))


def _make_figure_docx(tmp_path: Path, caption_text: str) -> Path:
    document = Document()
    document.add_paragraph("Текст перед рисунком.")
    _add_figure(document)
    document.add_paragraph(caption_text)
    document.add_paragraph("Текст после рисунка.")
    path = tmp_path / "input.docx"
    document.save(path)
    return path


def test_figure_caption_keeps_original_number_by_default(tmp_path: Path, gost_profile: Profile) -> None:
    # GOST-профиль по умолчанию: renumber=false (решение из плана — иначе
    # ломаются ручные ссылки вида «см. рисунок 3» в тексте отчёта).
    input_path = _make_figure_docx(tmp_path, "Рисунок 5 - Внешний вид установки")

    out_path, report = Formatter(gost_profile).run(input_path)

    document = Document(out_path)
    caption = document.paragraphs[2]
    assert caption.text == "Рисунок 5 — Внешний вид установки"
    assert caption.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER

    assert any(c.rule_id == "captions" for c in report.changes)


def test_figure_caption_renumbered_when_enabled(tmp_path: Path, gost_profile: Profile) -> None:
    profile = replace(
        gost_profile,
        captions=replace(
            gost_profile.captions, figure=replace(gost_profile.captions.figure, renumber=True)
        ),
    )
    document = Document()
    for intro_text, caption_text in (
        ("Текст перед первым рисунком.", "Рисунок 7 — Первая схема"),
        ("Текст перед вторым рисунком.", "Рисунок 9 — Вторая схема"),
    ):
        document.add_paragraph(intro_text)
        _add_figure(document)
        document.add_paragraph(caption_text)
    path = tmp_path / "input.docx"
    document.save(path)

    out_path, _ = Formatter(profile).run(path)

    result = Document(out_path)
    captions = [p.text for p in result.paragraphs if p.text.startswith("Рисунок")]
    assert captions == ["Рисунок 1 — Первая схема", "Рисунок 2 — Вторая схема"]


def test_table_caption_above_is_formatted(tmp_path: Path, gost_profile: Profile) -> None:
    document = Document()
    document.add_paragraph("Таблица 1 - Исходные данные")
    document.add_table(rows=1, cols=1)
    path = tmp_path / "input.docx"
    document.save(path)

    out_path, report = Formatter(gost_profile).run(path)

    result = Document(out_path)
    caption = result.paragraphs[0]
    assert caption.text == "Таблица 1 — Исходные данные"
    assert caption.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT

    assert any(c.rule_id == "captions" for c in report.changes)


def test_idempotent_with_figure_caption(tmp_path: Path, gost_profile: Profile) -> None:
    input_path = _make_figure_docx(tmp_path, "Рисунок 5 - Внешний вид установки")

    first_out, first_report = Formatter(gost_profile).run(input_path)
    assert any(c.rule_id == "captions" for c in first_report.changes)

    _, second_report = Formatter(gost_profile).run(first_out)

    assert second_report.changes == []
