from __future__ import annotations

from dataclasses import replace
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_formatter.core.oxml import paragraph_has_page_field
from docx_formatter.core.pipeline import Formatter
from docx_formatter.core.profile import Profile


def _make_docx(tmp_path: Path) -> Path:
    document = Document()
    document.add_paragraph("Первая страница.")
    document.add_page_break()
    document.add_paragraph("Вторая страница.")
    path = tmp_path / "input.docx"
    document.save(path)
    return path


def test_footer_gets_centered_page_field(tmp_path: Path, gost_profile: Profile) -> None:
    input_path = _make_docx(tmp_path)

    out_path, report = Formatter(gost_profile).run(input_path)

    document = Document(out_path)
    footer_paragraph = document.sections[0].footer.paragraphs[0]

    assert paragraph_has_page_field(footer_paragraph) is True
    assert footer_paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert any(c.rule_id == "page_numbers" for c in report.changes)


def test_first_page_has_no_page_field_when_skip_first_page(tmp_path: Path, gost_profile: Profile) -> None:
    input_path = _make_docx(tmp_path)

    out_path, _ = Formatter(gost_profile).run(input_path)

    document = Document(out_path)
    section = document.sections[0]
    assert section.different_first_page_header_footer is True
    first_footer_paragraph = section.first_page_footer.paragraphs[0]
    assert paragraph_has_page_field(first_footer_paragraph) is False


def test_applies_to_every_section(tmp_path: Path, gost_profile: Profile) -> None:
    document = Document()
    document.add_paragraph("Раздел 1.")
    document.add_section()
    document.add_paragraph("Раздел 2.")
    path = tmp_path / "input.docx"
    document.save(path)

    out_path, _ = Formatter(gost_profile).run(path)

    result = Document(out_path)
    assert len(result.sections) == 2
    for section in result.sections:
        assert paragraph_has_page_field(section.footer.paragraphs[0])


def test_disabled_in_profile_skips_rule(tmp_path: Path, gost_profile: Profile) -> None:
    profile = replace(gost_profile, page_numbers=replace(gost_profile.page_numbers, enabled=False))
    input_path = _make_docx(tmp_path)

    _, report = Formatter(profile).run(input_path)

    assert all(c.rule_id != "page_numbers" for c in report.changes)


def test_idempotent_with_page_numbers(tmp_path: Path, gost_profile: Profile) -> None:
    input_path = _make_docx(tmp_path)

    first_out, first_report = Formatter(gost_profile).run(input_path)
    assert any(c.rule_id == "page_numbers" for c in first_report.changes)

    _, second_report = Formatter(gost_profile).run(first_out)

    assert second_report.changes == []
