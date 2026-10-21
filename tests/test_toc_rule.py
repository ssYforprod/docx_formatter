from __future__ import annotations

from dataclasses import replace
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from docx_formatter.core.oxml import paragraph_has_toc_field
from docx_formatter.core.pipeline import Formatter
from docx_formatter.core.profile import Profile


def _make_docx(tmp_path: Path) -> Path:
    document = Document()
    document.add_paragraph("Текст отчёта.")
    path = tmp_path / "input.docx"
    document.save(path)
    return path


def test_disabled_by_default(tmp_path: Path, gost_profile: Profile) -> None:
    # Решение из плана: python-docx не умеет сам посчитать номера
    # страниц, поэтому TOC выключен в GOST-профиле по умолчанию.
    assert gost_profile.toc.enabled is False
    input_path = _make_docx(tmp_path)

    _, report = Formatter(gost_profile).run(input_path)

    assert all(c.rule_id != "toc" for c in report.changes)


def test_inserts_toc_field_at_start_when_enabled(tmp_path: Path, gost_profile: Profile) -> None:
    profile = replace(gost_profile, toc=replace(gost_profile.toc, enabled=True))
    input_path = _make_docx(tmp_path)

    out_path, report = Formatter(profile).run(input_path)

    document = Document(out_path)
    assert paragraph_has_toc_field(document.paragraphs[0]) is True
    assert any(c.rule_id == "toc" for c in report.changes)


def test_enables_update_fields_on_open(tmp_path: Path, gost_profile: Profile) -> None:
    profile = replace(gost_profile, toc=replace(gost_profile.toc, enabled=True))
    input_path = _make_docx(tmp_path)

    out_path, _ = Formatter(profile).run(input_path)

    document = Document(out_path)
    update_fields = document.settings.element.find(qn("w:updateFields"))
    assert update_fields is not None
    assert update_fields.get(qn("w:val")) == "true"


def test_idempotent_with_toc(tmp_path: Path, gost_profile: Profile) -> None:
    profile = replace(gost_profile, toc=replace(gost_profile.toc, enabled=True))
    input_path = _make_docx(tmp_path)

    first_out, first_report = Formatter(profile).run(input_path)
    assert any(c.rule_id == "toc" for c in first_report.changes)

    second_out, second_report = Formatter(profile).run(first_out)

    assert second_report.changes == []
    document = Document(second_out)
    toc_paragraphs = [p for p in document.paragraphs if paragraph_has_toc_field(p)]
    assert len(toc_paragraphs) == 1
