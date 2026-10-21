from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from factories import make_docx

from docx_formatter.core.pipeline import Formatter, output_path_for
from docx_formatter.core.profile import Profile


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_output_path_naming() -> None:
    assert output_path_for(Path("report.docx")) == Path("report_formatted.docx")


def test_original_file_is_never_modified_and_copy_is_created(
    tmp_path: Path, gost_profile: Profile
) -> None:
    input_path = make_docx(tmp_path, blocks=[("p", "Привет, мир")])
    before = _sha256(input_path)

    out_path, _ = Formatter(gost_profile).run(input_path)

    assert out_path == tmp_path / "input_formatted.docx"
    assert out_path.exists()
    assert _sha256(input_path) == before


def test_body_font_covers_all_scripts(tmp_path: Path, gost_profile: Profile) -> None:
    input_path = make_docx(tmp_path, blocks=[("p", "Привет, мир — hello world")])

    out_path, _ = Formatter(gost_profile).run(input_path)

    document = Document(out_path)
    paragraph = document.paragraphs[0]
    assert len(paragraph.runs) >= 1
    for run in paragraph.runs:
        rPr = run._element.find(qn("w:rPr"))
        rFonts = rPr.find(qn("w:rFonts"))
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            assert rFonts.get(qn(attr)) == "Times New Roman"
        for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            assert rFonts.get(qn(theme_attr)) is None


def test_line_spacing_and_indent(tmp_path: Path, gost_profile: Profile) -> None:
    input_path = make_docx(tmp_path, blocks=[("p", "Текст абзаца")])

    out_path, _ = Formatter(gost_profile).run(input_path)

    document = Document(out_path)
    paragraph = document.paragraphs[0]
    pf = paragraph.paragraph_format
    assert pf.line_spacing == 1.5

    pPr = paragraph._element.find(qn("w:pPr"))
    spacing = pPr.find(qn("w:spacing"))
    assert spacing.get(qn("w:line")) == "360"
    assert spacing.get(qn("w:lineRule")) == "auto"
    ind = pPr.find(qn("w:ind"))
    assert ind.get(qn("w:firstLine")) == "709"


def test_custom_output_path(tmp_path: Path, gost_profile: Profile) -> None:
    input_path = make_docx(tmp_path, blocks=[("p", "Текст")])
    custom_out = tmp_path / "куда_хочу.docx"

    out_path, _ = Formatter(gost_profile).run(input_path, output_path=custom_out)

    assert out_path == custom_out
    assert custom_out.exists()


def test_idempotent(tmp_path: Path, gost_profile: Profile) -> None:
    input_path = make_docx(
        tmp_path, blocks=[("p", "Первый абзац"), ("p", "Второй абзац")]
    )

    first_out, first_report = Formatter(gost_profile).run(input_path)
    assert first_report.changes != []

    second_out, second_report = Formatter(gost_profile).run(first_out)

    assert second_report.changes == []
    assert second_out.exists()
