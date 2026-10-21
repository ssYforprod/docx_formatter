from __future__ import annotations

from dataclasses import replace
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from docx_formatter.core.pipeline import Formatter
from docx_formatter.core.profile import Profile


def _save(document: Document, tmp_path: Path, name: str = "input.docx") -> Path:
    path = tmp_path / name
    document.save(path)
    return path


def _make_manual_heading_docx(tmp_path: Path) -> Path:
    document = Document()
    run = document.add_paragraph().add_run("1 ВВЕДЕНИЕ")
    run.bold = True
    run.font.size = Pt(16)
    document.add_paragraph("Текст основной части отчёта после заголовка.")
    return _save(document, tmp_path)


def test_manual_heading_detected_and_formatted(tmp_path: Path, gost_profile: Profile) -> None:
    input_path = _make_manual_heading_docx(tmp_path)

    out_path, report = Formatter(gost_profile).run(input_path)

    document = Document(out_path)
    heading = document.paragraphs[0]
    level1 = next(level for level in gost_profile.headings.levels if level.level == 1)

    assert heading.paragraph_format.page_break_before is True
    assert heading.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    for run in heading.runs:
        assert run.bold is level1.bold
        assert run.font.size == Pt(level1.size_pt)

    heading_changes = [c for c in report.changes if c.rule_id == "headings"]
    assert any(c.level == "info" and "заголовок" in c.what for c in heading_changes)


def test_ambiguous_paragraph_is_left_untouched_but_warned(
    tmp_path: Path, gost_profile: Profile
) -> None:
    document = Document()
    run = document.add_paragraph().add_run("Общие положения")
    run.bold = True
    path = _save(document, tmp_path)

    out_path, report = Formatter(gost_profile).run(path)

    result_doc = Document(out_path)
    paragraph = result_doc.paragraphs[0]
    # Абзац остался оформлен как body_text (выравнивание по ширине из
    # профиля, без разрыва страницы) — headings сознательно его не тронул,
    # раз уверенность пограничная.
    assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert paragraph.paragraph_format.page_break_before is not True

    heading_changes = [c for c in report.changes if c.rule_id == "headings"]
    assert any(c.level == "warning" for c in heading_changes)


def test_trust_styles_only_skips_heuristic_headings(tmp_path: Path, gost_profile: Profile) -> None:
    profile = replace(gost_profile, headings=replace(gost_profile.headings, trust_styles_only=True))
    document = Document()
    run = document.add_paragraph().add_run("1 ВВЕДЕНИЕ")
    run.bold = True
    run.font.size = Pt(16)
    path = _save(document, tmp_path)

    _, report = Formatter(profile).run(path)

    heading_changes = [c for c in report.changes if c.rule_id == "headings"]
    assert heading_changes == []


def test_idempotent_with_manual_heading(tmp_path: Path, gost_profile: Profile) -> None:
    input_path = _make_manual_heading_docx(tmp_path)

    first_out, first_report = Formatter(gost_profile).run(input_path)
    assert any(c.rule_id == "headings" for c in first_report.changes)

    _, second_report = Formatter(gost_profile).run(first_out)

    assert second_report.changes == []
