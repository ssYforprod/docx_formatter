from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.section import _Footer
from docx.shared import Pt
from docx.text.run import Run

from docx_formatter.core.context import DocContext
from docx_formatter.core.oxml import add_page_number_field, paragraph_has_page_field, set_run_font
from docx_formatter.core.profile import PageNumbersProfile, Profile
from docx_formatter.core.report import Change
from docx_formatter.core.units import lengths_equal, pt

# Единственное поддерживаемое положение номера — профиль (check в
# profile.py) допускает только "footer_center" для MVP.
_POSITION_ALIGNMENTS = {"footer_center": WD_ALIGN_PARAGRAPH.CENTER}

_FONT_ATTRS = ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia")


def _run_matches_font(run: Run, font_name: str, size: Pt) -> bool:
    if run.font.size is None or not lengths_equal(run.font.size, size):
        return False
    rPr = run._element.find(qn("w:rPr"))
    if rPr is None:
        return False
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        return False
    return all(rFonts.get(qn(attr)) == font_name for attr in _FONT_ATTRS)


def _footer_matches(footer: _Footer, cfg: PageNumbersProfile) -> bool:
    if footer.is_linked_to_previous:
        return False
    paragraphs = footer.paragraphs
    if len(paragraphs) != 1:
        return False
    paragraph = paragraphs[0]
    if paragraph.paragraph_format.alignment != _POSITION_ALIGNMENTS[cfg.position]:
        return False
    if not paragraph_has_page_field(paragraph):
        return False
    size = pt(cfg.size_pt)
    return all(_run_matches_font(run, cfg.font, size) for run in paragraph.runs)


def _rebuild_footer(footer: _Footer, cfg: PageNumbersProfile) -> None:
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.paragraph_format.alignment = _POSITION_ALIGNMENTS[cfg.position]

    add_page_number_field(paragraph)
    size = pt(cfg.size_pt)
    for run in paragraph.runs:
        set_run_font(run, cfg.font)
        run.font.size = size


def _first_page_footer_is_empty(footer: _Footer) -> bool:
    if footer.is_linked_to_previous:
        return False
    return all(not paragraph.text.strip() and not paragraph_has_page_field(paragraph) for paragraph in footer.paragraphs)


def _clear_first_page_footer(footer: _Footer) -> None:
    footer.is_linked_to_previous = False
    for paragraph in footer.paragraphs:
        for run in list(paragraph.runs):
            run._element.getparent().remove(run._element)


class PageNumbersRule:
    """Номер страницы в футере — центрированное поле PAGE на каждой
    секции документа (не только на первой — иначе в документах с
    несколькими разделами нумерация проставится лишь частично).

    При page_numbers.skip_first_page титульная страница получает
    отдельный, пустой first_page_footer: по умолчанию он связан с
    обычным футером (is_linked_to_previous=True) даже после включения
    different_first_page_header_footer — отвязывать и чистить нужно
    явно, иначе номер всё равно появится на титульном листе.
    """

    id = "page_numbers"

    def enabled_for(self, profile: Profile) -> bool:
        return profile.page_numbers.enabled

    def apply(self, context: DocContext) -> list[Change]:
        cfg = context.profile.page_numbers
        changes: list[Change] = []

        for index, section in enumerate(context.document.sections):
            changed = False

            if not _footer_matches(section.footer, cfg):
                _rebuild_footer(section.footer, cfg)
                changed = True

            if cfg.skip_first_page:
                if not section.different_first_page_header_footer:
                    section.different_first_page_header_footer = True
                    changed = True
                if not _first_page_footer_is_empty(section.first_page_footer):
                    _clear_first_page_footer(section.first_page_footer)
                    changed = True
            elif section.different_first_page_header_footer:
                section.different_first_page_header_footer = False
                changed = True

            if changed:
                changes.append(
                    Change(self.id, "info", f"раздел #{index + 1}", "настроена нумерация страниц")
                )

        return changes
